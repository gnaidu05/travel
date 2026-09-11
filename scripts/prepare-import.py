import json, sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

src=Path(sys.argv[1]); out=Path(sys.argv[2]); wb=load_workbook(src,read_only=True,data_only=True)
def clean(v):
    if isinstance(v,(datetime,date)): return v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
    if v is None: return None
    return str(v).strip() if isinstance(v,str) else v
def date_only(v):
    v=clean(v)
    return v if isinstance(v,str) and len(v)==10 and v[4]=='-' and v[7]=='-' else None
team=[]; assignments=[]; ws=wb['Team']; headers=[clean(x.value) for x in next(ws.iter_rows())]
for source_row,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
    if not row[0]: continue
    email=str(row[1] or '').strip().lower()
    member={'source_row':source_row,'full_name':clean(row[0]),'email':email or None,'city':clean(row[2]),'unified_grade':clean(row[3]),'category':clean(row[4]),'can_travel':str(row[5]).lower()=='yes'}
    team.append(member)
    for i,val in enumerate(row[6:],6):
        if val not in (None,''):
            raw=clean(val); low=str(raw).lower(); status='Travel' if 'travel' in low else 'Return' if 'return' in low else 'Unavailable' if 'unavailable' in low else 'Assessment'
            assignments.append({'member_source_row':source_row,'assignment_date':headers[i],'status':status,'college_label':raw})
ws=wb['Overall']; headers=[clean(x.value) for x in next(ws.iter_rows())]; colleges=[]
def val(row,n): return clean(row[headers.index(n)]) if n in headers else None
for num,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
    if not row[0]: continue
    source={str(headers[i]):clean(v) for i,v in enumerate(row) if v not in (None,'')}
    colleges.append({'source_row':num,'college_name':val(row,'College name'),'college_group':val(row,'College Group(CG)'),'college_city':val(row,'College City'),'college_state':val(row,'College State'),'college_zone':val(row,'College Zone'),'category':val(row,'Category'),'college_type':val(row,'Type'),'nirf_2023':val(row,'NIRF 2023'),'nirf_2024':val(row,'NIRF 2024'),'nirf_2025':val(row,'NIRF 2025'),'salary_bands':val(row,'Salaries to be Sent'),'hiring_type':val(row,'Type of Hiring'),'assessment_framework_1':val(row,'Assessment-1'),'assessment_framework_2':val(row,'Assessment-2'),'assessment_framework_3':val(row,'Assessment-3'),'registration_timeline':val(row,'Registration Timeline'),'hiring_date':val(row,'Hiring'),'application_date':val(row,'Application'),'application_invited':val(row,'Application Invited'),'jd_template':val(row,'JD Template'),'bu_identified':val(row,'BU Identified'),'registration_date':date_only(val(row,'Registration')),'english_date':date_only(val(row,'English')),'assessment_date':date_only(val(row,'Physical Assessment-1')),'assessment_date_2':date_only(val(row,'Physical Assessment-2')),'assessment_date_3':date_only(val(row,'Physical Assessment-3')),'interview_date':date_only(val(row,'Interview Date')),'nearest_airport':val(row,'Nearest Airport'),'airport_distance_km':val(row,'Distance'),'nearest_station':val(row,'Nearest Station'),'station_distance_km':clean(row[41]),'source_data':source})
out.parent.mkdir(parents=True,exist_ok=True); out.write_text(json.dumps({'team_members':team,'colleges':colleges,'assignments':assignments},ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(f'Prepared {len(team)} team members, {len(colleges)} colleges and {len(assignments)} dated roster entries.')
